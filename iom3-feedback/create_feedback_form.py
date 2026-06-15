from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = "/Users/gillian/Documents/New project 2/iom3-feedback/Registration Application Feedback Form - Completed.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_question(doc, number, question, answer, comment=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(f"{number}. {question}")
    run.bold = True

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.8)
    labels = ["Response", "Comments"]
    values = [answer, comment or "N/A"]

    for i in range(2):
      cell = table.rows[0].cells[i]
      set_cell_shading(cell, "DCE6F1")
      p = cell.paragraphs[0]
      p.alignment = WD_ALIGN_PARAGRAPH.CENTER
      r = p.add_run(labels[i])
      r.bold = True

    row = table.add_row().cells
    row[0].text = answer
    row[1].text = comment or "N/A"
    doc.add_paragraph()


def add_overall_satisfaction(doc, selected):
    p = doc.add_paragraph()
    run = p.add_run("11. Overall, how satisfied or dissatisfied are you with the process?")
    run.bold = True

    options = [
        "Very satisfied",
        "Somewhat satisfied",
        "Neither",
        "Somewhat dissatisfied",
        "Very dissatisfied",
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_cell_shading(table.rows[0].cells[0], "DCE6F1")
    set_cell_shading(table.rows[0].cells[1], "DCE6F1")
    table.rows[0].cells[0].paragraphs[0].add_run("Option").bold = True
    table.rows[0].cells[1].paragraphs[0].add_run("Selected").bold = True
    for option in options:
        row = table.add_row().cells
        row[0].text = option
        row[1].text = "Yes" if option == selected else ""
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("INSTITUTE OF MATERIALS, MINERALS AND MINING")
r.bold = True
r.font.size = Pt(15)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Registration Application Feedback Form")
r.bold = True
r.font.size = Pt(13)

intro = doc.add_paragraph(
    "Following your recent professional review interview with IOM3, please find my completed feedback form below."
)
intro.alignment = WD_ALIGN_PARAGRAPH.LEFT

add_question(doc, 1, "Did you attend a workshop?", "Yes")
add_question(doc, 2, "Are there any areas that you felt weren’t covered in the workshop?", "No", "The workshop covered the key points clearly and helped me understand the process well.")
add_question(doc, 3, "Did you receive Mentor support to assist you with the application documents / process?", "Yes", "My mentor was not from my organisation, but the support was helpful and valuable throughout the process.")
add_question(doc, 4, "Were you given sufficient briefing information on and notice of your Professional Review Interview?", "Yes")
add_question(doc, 5, "Do you feel that the Interview Panel members had fully prepared themselves on your application prior to the interview?", "Yes")
add_question(doc, 6, "Did the Interview Panel give you every opportunity to respond fully to the questions they put to you?", "Yes")
add_question(doc, 7, "Were there any areas that you felt that you weren’t able to cover?", "No", "I felt I had the opportunity to explain my experience and respond to the questions raised.")
add_question(doc, 8, "Did you feel that the interview panel conducted themselves in a professional manner?", "Yes", "The interviewers were very professional, polite, and made the interview feel fair and well organised.")
add_question(doc, 9, "Did you encounter any issues with the remote interview?", "No", "The remote interview process worked smoothly and I did not encounter any technical issues.")
add_question(doc, 10, "Would you be interested in acting as an IOM3 Volunteer in the future?", "Yes")
add_overall_satisfaction(doc, "Very satisfied")

p = doc.add_paragraph()
run = p.add_run("12. Do you have any other comments or suggestions?")
run.bold = True
doc.add_paragraph(
    "Overall, I was very satisfied with the process. The information provided in advance was helpful, and the interview panel were professional, polite, and supportive throughout the interview."
)

doc.add_paragraph()
meta = doc.add_table(rows=3, cols=2)
meta.style = "Table Grid"
labels = ["Name", "Membership Number", "Date of Interview"]
values = ["Dr Xingran (Gillian) Gao", "711359", "18 May 2026"]
for i, (label, value) in enumerate(zip(labels, values)):
    set_cell_shading(meta.rows[i].cells[0], "DCE6F1")
    meta.rows[i].cells[0].paragraphs[0].add_run(label).bold = True
    meta.rows[i].cells[1].text = value

doc.save(OUTPUT)
print(OUTPUT)
